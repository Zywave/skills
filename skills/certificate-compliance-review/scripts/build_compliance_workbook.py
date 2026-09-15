#!/usr/bin/env python3
"""
Build a COI compliance review workbook from coi_data.json.

Usage:
    python3 build_compliance_workbook.py <coi_data.json> <output.xlsx> [--no-logo]

Tabs produced:
    Summary            verdict, counts, blockers, what the certificate does evidence
    Compliance Matrix  one row per requirement, formula-driven status
    Source Data        every extracted value with the document it came from

Status is a formula, never a hardcoded string. Numeric rows compare Req Amt to Actual Amt;
qualitative rows read the blue "Met?" input column. That way the workbook recalculates if
someone corrects a limit, and a reviewer can see exactly why each row landed where it did.
"""

import argparse, json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

LOGO_CANDIDATES = [
    "/mnt/skills/organization/brand-applicator/assets/logos/"
    "Logo_Zywave_without-tagline_white and green.png",
]

NAVY, MOON, GRAY, WHITE = "051C2C", "0F2D52", "F7F7F7", "FFFFFF"
OK_F = PatternFill("solid", fgColor="D8F0CB")
BAD_F = PatternFill("solid", fgColor="F8D0D0")
PART_F = PatternFill("solid", fgColor="FDEBC8")

H1 = Font(name="Arial", bold=True, size=16, color=WHITE)
H2 = Font(name="Arial", bold=True, size=10, color=WHITE)
LBL = Font(name="Arial", bold=True, size=9, color=NAVY)
BODY = Font(name="Calibri", size=10, color="1A1A1A")
BODY_B = Font(name="Calibri", bold=True, size=10, color=NAVY)
SMALL = Font(name="Calibri", size=8, italic=True, color="5A5A5A")
THIN = Side(style="thin", color="C8C8C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True)

HDR = ["ID", "Coverage", "Requirement", "Contract Ref", "Required", "Req Amt",
       "Evidence Source", "Actual / Finding", "Actual Amt", "Met?", "Status",
       "Severity", "Action Required"]
WID = [8, 20, 38, 12, 26, 13, 24, 46, 13, 8, 16, 10, 44]
HR = 4  # header row of the matrix


def g(d, *path, default=""):
    cur = d
    for p in path:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(p)
        if cur is None:
            return default
    return cur


def build_matrix(wb, d):
    reqs = g(d, "contract", "requirements", default=[]) or []
    if not reqs:
        sys.exit("coi_data.json has no contract.requirements — nothing to review.")

    ws = wb.create_sheet("Compliance Matrix")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:M1")
    ws["A1"] = "COI COMPLIANCE MATRIX — CONTRACT REQUIREMENTS vs. POLICIES IN FORCE"
    ws["A1"].font = H1
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 30

    sub = (f"Insured: {g(d,'insured','name')}   |   "
           f"Certificate Holder: {g(d,'certificate_holder','name')}   |   "
           f"Agreement {g(d,'contract','agreement_number')}   |   "
           f"Reviewed {g(d,'certificate','issue_date')}")
    if g(d, "meta", "specimen"):
        sub += "   |   MOCK DEMONSTRATION DATA"
    ws.merge_cells("A2:M2")
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", size=9, color=WHITE)
    ws["A2"].fill = PatternFill("solid", fgColor=MOON)
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 18

    for i, h in enumerate(HDR, start=1):
        c = ws.cell(row=HR, column=i, value=h)
        c.font, c.border = H2, BOX
        c.fill = PatternFill("solid", fgColor=MOON)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[HR].height = 26
    for i, w in enumerate(WID, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = HR + 1
    for q in reqs:
        ws.cell(row=r, column=1, value=q.get("id", ""))
        ws.cell(row=r, column=2, value=q.get("coverage", ""))
        ws.cell(row=r, column=3, value=q.get("requirement", ""))
        ws.cell(row=r, column=4, value=q.get("contract_ref", ""))
        ws.cell(row=r, column=5, value=q.get("required_text", ""))
        ws.cell(row=r, column=6, value=q.get("required_amount"))
        ws.cell(row=r, column=7, value=q.get("evidence_source", ""))
        ws.cell(row=r, column=8, value=q.get("actual_text", ""))
        ws.cell(row=r, column=9, value=q.get("actual_amount"))
        ws.cell(row=r, column=10, value=q.get("met", ""))
        ws.cell(row=r, column=11, value=(
            f'=IF(F{r}="",'
            f'IF(J{r}="Y","COMPLIANT",IF(J{r}="PARTIAL","PARTIAL","NOT COMPLIANT")),'
            f'IF(I{r}="","NOT COMPLIANT",IF(I{r}>=F{r},"COMPLIANT","NOT COMPLIANT")))'))
        ws.cell(row=r, column=12, value=q.get("severity", ""))
        ws.cell(row=r, column=13, value=q.get("action", ""))

        for col in range(1, 14):
            c = ws.cell(row=r, column=col)
            c.border, c.alignment, c.font = BOX, TOP, BODY
        ws.cell(row=r, column=1).font = BODY_B
        ws.cell(row=r, column=10).font = Font(name="Calibri", size=10, bold=True, color="0000FF")
        ws.cell(row=r, column=10).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=11).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=r, column=11).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=12).alignment = Alignment(horizontal="center", vertical="top")
        for col in (6, 9):
            ws.cell(row=r, column=col).number_format = '$#,##0;($#,##0);-'
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="right", vertical="top")
        r += 1

    last = r - 1
    rng = f"K{HR+1}:K{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$K{HR+1}="COMPLIANT"'], fill=OK_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$K{HR+1}="NOT COMPLIANT"'], fill=BAD_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$K{HR+1}="PARTIAL"'], fill=PART_F))
    sv = f"L{HR+1}:L{last}"
    ws.conditional_formatting.add(sv, FormulaRule(formula=[f'$L{HR+1}="Critical"'], fill=BAD_F))
    ws.conditional_formatting.add(sv, FormulaRule(formula=[f'$L{HR+1}="High"'], fill=PART_F))

    dv = DataValidation(type="list", formula1='"Y,N,PARTIAL"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J{HR+1}:J{last}")

    ws.print_area = f"A1:M{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"
    ws.auto_filter.ref = f"A{HR}:M{last}"
    ws.freeze_panes = f"A{HR+1}"

    n = last + 2
    ws.cell(row=n, column=1, value=(
        "Column J (blue) is the reviewer input for requirements that are not a numeric limit "
        "comparison. Column K derives from J, or from the Req Amt / Actual Amt comparison where "
        "both are populated. Do not overwrite column K.")).font = SMALL
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=13)
    if g(d, "meta", "specimen"):
        ws.cell(row=n + 1, column=1, value=(
            "All entities, carriers, policy numbers and ratings are fictitious. Prepared for "
            "software demonstration only; not a coverage opinion.")).font = SMALL
        ws.merge_cells(start_row=n + 1, start_column=1, end_row=n + 1, end_column=13)
    return last


def build_summary(wb, d, last, use_logo=True):
    sm = wb.create_sheet("Summary", 0)
    sm.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [3, 30, 26, 20, 20, 20, 20]):
        sm.column_dimensions[col].width = w

    for rr in range(1, 4):
        for cc in range(1, 8):
            sm.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=NAVY)
    if use_logo:
        for path in LOGO_CANDIDATES:
            if os.path.exists(path):
                try:
                    from openpyxl.drawing.image import Image as XLImage
                    img = XLImage(path)
                    ar = img.width / img.height
                    img.height = 34
                    img.width = int(34 * ar)
                    img.anchor = "B2"
                    sm.add_image(img)
                except Exception:
                    pass
                break
    sm.cell(row=2, column=4, value="CERTIFICATE OF INSURANCE — COMPLIANCE REVIEW").font = \
        Font(name="Arial", bold=True, size=13, color=WHITE)
    sm.cell(row=2, column=4).alignment = Alignment(vertical="center")
    sm.row_dimensions[2].height = 26

    req = g(d, "contract", "request", default={}) or {}
    rows = [
        ("Insured", g(d, "insured", "name")),
        ("Certificate Holder", g(d, "certificate_holder", "name")),
        ("Agreement", f'{g(d,"contract","agreement_number")} — {g(d,"contract","project")}'.strip(" —")),
        ("Project Site", g(d, "contract", "project_address")),
        ("Contract Term", f'{g(d,"contract","term_start")} – {g(d,"contract","term_end")}'.strip(" –")),
        ("Producer", g(d, "producer", "name")),
        ("Certificate Number", g(d, "certificate", "number")),
        ("Certificate Date", g(d, "certificate", "issue_date")),
        ("Holder Deadline", req.get("deadline", "")),
        ("Mobilization Date", req.get("mobilization_date", "")),
    ]
    row = 5
    for label, val in rows:
        if not val:
            continue
        sm.cell(row=row, column=2, value=label).font = LBL
        c = sm.cell(row=row, column=3, value=val)
        c.font, c.alignment = BODY, Alignment(horizontal="left")
        sm.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        row += 1

    row += 1
    sm.cell(row=row, column=2, value="REVIEW RESULT").font = Font(
        name="Arial", bold=True, size=11, color=NAVY)
    row += 1
    head = row
    for i, h in enumerate(["Status", "Count", "% of Requirements"], start=2):
        c = sm.cell(row=head, column=i, value=h)
        c.font, c.border = H2, BOX
        c.fill = PatternFill("solid", fgColor=MOON)
        c.alignment = Alignment(horizontal="center", vertical="center")
    sm.row_dimensions[head].height = 20

    M = "'Compliance Matrix'"
    stat = f"{M}!$K${HR+1}:$K${last}"
    row = head + 1
    for label in ["COMPLIANT", "PARTIAL", "NOT COMPLIANT"]:
        sm.cell(row=row, column=2, value=label).font = BODY_B
        sm.cell(row=row, column=3, value=f'=COUNTIF({stat},B{row})').font = BODY
        row += 1
    tot = row
    sm.cell(row=tot, column=2, value="TOTAL REQUIREMENTS").font = BODY_B
    sm.cell(row=tot, column=3, value=f'=COUNTA({M}!$A${HR+1}:$A${last})').font = BODY_B
    for rr in range(head + 1, tot):
        c = sm.cell(row=rr, column=4, value=f'=IFERROR(C{rr}/$C${tot},0)')
        c.number_format, c.font = "0.0%", BODY
    for rr in range(head + 1, tot + 1):
        for cc in range(2, 5):
            sm.cell(row=rr, column=cc).border = BOX
            sm.cell(row=rr, column=cc).alignment = Alignment(horizontal="center")
        sm.cell(row=rr, column=2).alignment = Alignment(horizontal="left", indent=1)
    sm.conditional_formatting.add(f"B{head+1}:D{head+1}", FormulaRule(formula=["TRUE"], fill=OK_F))
    sm.conditional_formatting.add(f"B{head+2}:D{head+2}", FormulaRule(formula=["TRUE"], fill=PART_F))
    sm.conditional_formatting.add(f"B{head+3}:D{head+3}", FormulaRule(formula=["TRUE"], fill=BAD_F))

    row = tot + 2
    for label, crit in [("Critical findings", "Critical"), ("High findings", "High")]:
        sm.cell(row=row, column=2, value=label).font = LBL
        c = sm.cell(row=row, column=3,
                    value=f'=COUNTIFS({M}!$L${HR+1}:$L${last},"{crit}")')
        c.font, c.alignment = BODY_B, Alignment(horizontal="center")
        row += 1
    sm.cell(row=row, column=2, value="Overall determination").font = LBL
    c = sm.cell(row=row, column=3, value=(
        f'=IF(COUNTIF({stat},"NOT COMPLIANT")>0,'
        f'"NOT COMPLIANT — CERTIFICATE DOES NOT SATISFY CONTRACT REQUIREMENTS",'
        f'IF(COUNTIF({stat},"PARTIAL")>0,"CONDITIONALLY COMPLIANT","COMPLIANT"))'))
    c.font = Font(name="Arial", bold=True, size=10, color="9C0006")
    sm.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)

    blockers = [q for q in g(d, "contract", "requirements", default=[])
                if q.get("severity") in ("Critical", "High") and q.get("action")]
    seen, uniq = set(), []
    for b in blockers:
        key = b.get("action", "")[:60]
        if key not in seen:
            seen.add(key)
            uniq.append(b)
    if uniq:
        row += 2
        sm.cell(row=row, column=2, value="BLOCKERS TO ISSUING A COMPLIANT CERTIFICATE").font = \
            Font(name="Arial", bold=True, size=11, color=NAVY)
        row += 1
        for n, b in enumerate(uniq, start=1):
            sm.cell(row=row, column=2, value=n).font = Font(
                name="Arial", bold=True, size=10, color="9C0006")
            sm.cell(row=row, column=2).alignment = Alignment(horizontal="center", vertical="top")
            txt = f'{b.get("coverage","")} — {b.get("requirement","")}: {b.get("action","")} ({b.get("id","")})'
            c = sm.cell(row=row, column=3, value=txt)
            c.font, c.alignment = BODY, TOP
            sm.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
            sm.row_dimensions[row].height = 30
            row += 1

    row += 2
    if g(d, "meta", "specimen"):
        sm.cell(row=row, column=2, value=(
            "Mock demonstration data. All entities, carriers, policy numbers, NAIC codes and "
            "A.M. Best ratings are fictitious.")).font = SMALL
    else:
        sm.cell(row=row, column=2, value=(
            "This workbook is a compliance comparison, not a coverage opinion and not legal "
            "advice. Coverage is determined by the policies themselves.")).font = SMALL
    sm.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)


def build_source(wb, d):
    sd = wb.create_sheet("Source Data")
    sd.sheet_view.showGridLines = False
    for col, w in zip("ABCD", [3, 34, 52, 34]):
        sd.column_dimensions[col].width = w
    sd.merge_cells("A1:D1")
    sd["A1"] = "EXTRACTED SOURCE DATA"
    sd["A1"].font = H1
    sd["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sd["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    sd.row_dimensions[1].height = 28

    sections = []
    ct = d.get("contract", {}) or {}
    items = [(k.replace("_", " ").title(), str(v), "Contract")
             for k, v in ct.items()
             if k not in ("requirements", "request") and v not in (None, "", [], {})]
    hold = d.get("certificate_holder", {}) or {}
    if hold.get("name"):
        items.append(("Certificate holder wording",
                      ", ".join(str(hold.get(x)) for x in
                                ("name", "address1", "address2", "city", "state", "zip")
                                if hold.get(x)),
                      hold.get("verbatim_source", "Contract")))
    if items:
        sections.append((f'Contract — {ct.get("agreement_number","")}'.strip(" —"), items))

    rq = ct.get("request") or {}
    if rq:
        sections.append(("Certificate request", [
            (k.replace("_", " ").title(),
             "; ".join(v) if isinstance(v, list) else str(v), "Request")
            for k, v in rq.items() if v not in (None, "", [], {})]))

    labels = {"general_liability": "General Liability", "automobile": "Business Auto",
              "umbrella": "Umbrella / Excess", "workers_comp": "Workers Compensation"}
    insurers = {i.get("letter"): i for i in d.get("insurers", [])}
    for key, label in labels.items():
        pol = (d.get("policies") or {}).get(key)
        if not pol:
            sections.append((f"{label} — no policy in force",
                             [("Status", "NO POLICY PROVIDED", "—")]))
            continue
        ins = insurers.get(pol.get("insurer_letter"), {})
        it = [("Carrier", f'{ins.get("name","")} (NAIC {ins.get("naic","")})', "Declarations"),
              ("A.M. Best", ins.get("am_best") or "Not stated", "Declarations"),
              ("Policy number", pol.get("policy_number", ""), "Declarations"),
              ("Policy period", f'{pol.get("effective","")} – {pol.get("expiration","")}',
               "Declarations")]
        for lk, lv in (pol.get("limits") or {}).items():
            if lv is not None:
                it.append((lk.replace("_", " ").title(), f"${lv:,}", "Limits"))
        for e in pol.get("endorsements") or []:
            it.append((f'{e.get("form","")} {e.get("edition","")}'.strip(),
                       f'{e.get("title","")}'
                       + (f' — {e.get("scope")}' if e.get("scope") else ""),
                       "Endorsement schedule"))
        sections.append((f'{label} — {pol.get("policy_number","")}', it))

    r = 3
    for title, items in sections:
        sd.cell(row=r, column=2, value=title).font = Font(
            name="Arial", bold=True, size=10, color=WHITE)
        for cc in range(2, 5):
            sd.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=MOON)
        r += 1
        for h, cc in zip(["Field", "Extracted Value", "Located In"], range(2, 5)):
            sd.cell(row=r, column=cc, value=h).font = LBL
            sd.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=GRAY)
            sd.cell(row=r, column=cc).border = BOX
        r += 1
        for a, b, cite in items:
            sd.cell(row=r, column=2, value=a).font = BODY_B
            sd.cell(row=r, column=3, value=b).font = BODY
            sd.cell(row=r, column=4, value=cite).font = BODY
            for cc in range(2, 5):
                sd.cell(row=r, column=cc).border = BOX
                sd.cell(row=r, column=cc).alignment = TOP
            r += 1
        r += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("out")
    ap.add_argument("--no-logo", action="store_true")
    a = ap.parse_args()

    d = json.load(open(a.data))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    last = build_matrix(wb, d)
    build_summary(wb, d, last, use_logo=not a.no_logo)
    build_source(wb, d)
    wb._sheets = [wb["Summary"], wb["Compliance Matrix"], wb["Source Data"]]
    wb.save(a.out)

    n = len(g(d, "contract", "requirements", default=[]))
    print(f"wrote {a.out}  ({n} requirements)")
    print("NEXT: run recalc.py — openpyxl writes formulas with no cached values, so the "
          "status column reads as blank until LibreOffice evaluates it.")


if __name__ == "__main__":
    main()
