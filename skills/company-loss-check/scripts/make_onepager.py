#!/usr/bin/env python3
"""Build a one-page, unbranded Word summary from a profile JSON.

Usage: python make_onepager.py profile.json onepager.docx [charts.png]
"""
import json
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor

FOOTER = ("Source: Zywave loss data. Publicly reported loss events, generally $1M or more. "
          "Not this company's own loss runs. Smaller companies are underrepresented, so a thin result "
          "is not evidence of a clean history. Lines of business show the coverage most likely to respond, "
          "not what paid.")


def money(n):
    if n is None:
        return "n/a"
    if n >= 1e9:
        return f"${n / 1e9:.2f}B"
    if n >= 1e6:
        return f"${n / 1e6:.1f}M"
    if n >= 1e3:
        return f"${n / 1e3:.0f}K"
    return f"${n:,.0f}"


def para(doc, text, size=10, bold=False, color=None, after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def main(src, out, chart=None):
    with open(src) as f:
        p = json.load(f)
    co, h = p["company"], p["headline"]
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Arial"

    para(doc, f"{co['name']}: publicly reported losses", size=16, bold=True, after=2)
    scope = "corporate family" if p.get("scope") == "family" else "entity"
    sub = f"{co.get('state', '')} | Ultimate parent: {co.get('ultimate_parent', co['name'])} | Scope: {scope} | Data pulled {p.get('pulled', '')}"
    para(doc, sub, size=9, color="6B6A66", after=8)

    family = p.get("scope") == "family"
    disclosure = f"{h['contributing']:,} of {h['cases']:,} records with an amount"
    if h.get("total_loss") is None or h.get("contributing", 0) == 0:
        total_txt = f"No dollar amount recorded ({disclosure})"
    else:
        total_txt = f"{money(h['total_loss'])} ({disclosure})"
    rows = [["Case records", f"{h['cases']:,}"], ["Total reported loss", total_txt]]
    if not family and h.get("average") is not None:
        rows.append(["Average (records with an amount)", money(h["average"])])
    if not family and h.get("largest") is not None:
        rows.append(["Largest single loss", money(h["largest"])])
    table(doc, ["Headline", "Value"], rows)
    para(doc, "", size=4, after=2)

    if family and p.get("entities"):
        table(doc, ["Entity", "State", "Case records", "Total reported loss", "With an amount"],
              [[r["name"], r.get("state", ""), f"{r['cases']:,}",
                money(r.get("total_loss")) if r.get("contributing", 0) > 0 else "No amount recorded",
                f"{r.get('contributing', 0):,} of {r['cases']:,}"] for r in p["entities"]])
        para(doc, "", size=4, after=2)
    else:
        lines = sorted(p.get("lines", []), key=lambda r: r.get("total_loss") or 0, reverse=True)[:6]
        if lines:
            table(doc, ["Line of business", "Case records", "Total reported loss", "With an amount"],
                  [[r["name"], r["cases"], money(r.get("total_loss")) if r.get("contributing", 0) > 0 else "No amount recorded",
                    f"{r.get('contributing', 0)} of {r['cases']}"] for r in lines])
            para(doc, "", size=4, after=2)
        if chart:
            doc.add_picture(chart, width=Inches(6.8))
        years = sorted(p.get("years", []), key=lambda r: r["year"], reverse=True)[:10]
        if years:
            table(doc, ["Year", "Case records", "With an amount"],
                  [[r["year"], r["cases"], r.get("contributing", 0)] for r in years])
            para(doc, "", size=4, after=2)

    for n in p.get("notes", [])[:4]:
        para(doc, "\u2022 " + n, size=9, after=1)
    para(doc, "", size=4, after=2)
    para(doc, FOOTER, size=7.5, color="6B6A66")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) == 4 else None)
