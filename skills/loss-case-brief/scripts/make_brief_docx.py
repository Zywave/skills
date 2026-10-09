#!/usr/bin/env python3
"""Build a one-page, unbranded Word brief from a profile JSON.

Usage: python make_brief_docx.py profile.json brief.docx
"""
import json
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor

FOOTER = ("Source: Zywave loss data. A publicly reported loss event, as recorded; wording follows the case's "
          "status and allegations are not findings. Not a legal opinion. Lines of business show the coverage "
          "most likely to respond, not what paid.")


def money(n):
    if n is None:
        return "n/a"
    if n >= 1e9:
        return f"${n / 1e9:.2f}B"
    if n >= 1e6:
        return f"${n / 1e6:.1f}M" if n < 1e8 else f"${n / 1e6:.0f}M"
    return f"${n / 1e3:.0f}K" if n >= 1e3 else f"${n:,.0f}"


def para(doc, text, size=10, bold=False, color=None, after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def bullets(doc, title, items, as_bullets=True):
    if not items:
        return
    para(doc, title, size=11, bold=True, after=2)
    for i in items:
        para(doc, ("\u2022 " if as_bullets else "") + i, size=9.5, after=1 if as_bullets else 4)


def main(src, out):
    p = json.load(open(src))
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    para(doc, p["title"], size=16, bold=True, after=2)
    parent = p.get("parent")
    sub = f"{p.get('company', '')}" + (f" | Parent: {parent}" if parent and parent != p.get("company") else "") + f" | Case {p.get('case_id', '')} | Data pulled {p.get('pulled', '')}"
    para(doc, sub, size=9, color="6B6A66", after=6)
    amt = p.get("amount")
    amt_txt = "No dollar amount recorded" if amt is None else money(amt) + (f" ({p['qualifier']})" if p.get("qualifier") else "")
    para(doc, f"Amount: {amt_txt}   |   Status: {p.get('status', '')}   |   Year: {p.get('year', '')}", size=10.5, bold=True, after=2)
    if amt is not None and p.get("nature"):
        para(doc, p["nature"], size=9, color="6B6A66", after=6)
    bullets(doc, "What happened", p.get("what_happened", []), as_bullets=False)
    cost = p.get("cost", [])
    if cost:
        para(doc, "What it cost", size=11, bold=True, after=2)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(("Component", "Amount", "Note")):
            t.rows[0].cells[i].text = h
        for c in cost:
            cells = t.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = c["label"], money(c.get("amount")), c.get("note", "")
        for row in t.rows:
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after = Pt(0)
                    for r in pp.runs:
                        r.font.size = Pt(9)
        para(doc, "The amount above is the all-in figure. Components are parts of it and are not added to it.", size=8.5, color="6B6A66", after=6)
    bullets(doc, "Where and who", p.get("where_who", []))
    bullets(doc, "Cause", p.get("cause", []))
    bullets(doc, "Related cases", p.get("related", []))
    bullets(doc, "Lines of business tagged", p.get("lines", []))
    bullets(doc, "Worth knowing", p.get("notes", [])[:4])
    para(doc, "", size=4, after=2)
    para(doc, FOOTER + f" Snapshot as of {p.get('pulled', '')}.", size=7.5, color="6B6A66")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
